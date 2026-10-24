"""Import and export helpers for batch stock valuation."""
from __future__ import annotations

import csv
import os
from dataclasses import dataclass
from typing import Optional

from valuation import ValuationResult


@dataclass
class BatchItem:
    symbol: str
    market: str
    quantity: Optional[float] = None
    cost_basis: Optional[float] = None
    target_weight: Optional[float] = None


_MARKETS = {
    "US": "US", "USA": "US", "NYSE": "US", "NASDAQ": "US", "AMEX": "US",
    "NSE": "NSE", "NSEI": "NSE", "INDIA NSE": "NSE",
    "BSE": "BSE", "BOM": "BSE", "INDIA BSE": "BSE",
}


def _market(value: object, symbol: str, default_market: str) -> str:
    raw = str(value or "").strip().upper()
    if raw in _MARKETS:
        return _MARKETS[raw]
    upper = symbol.upper()
    if upper.endswith(".NS"):
        return "NSE"
    if upper.endswith(".BO"):
        return "BSE"
    return default_market


def _rows_from_text(path: str) -> list[list[object]]:
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(4096)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [list(row) for row in csv.reader(handle, dialect) if any(str(c).strip() for c in row)]


def _rows_from_xlsx(path: str) -> list[list[object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Excel import requires openpyxl. Run: pip install -r requirements.txt") from exc
    book = load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True) if any(v is not None and str(v).strip() for v in row)]
    book.close()
    return rows


def import_stock_list(path: str, default_market: str = "US", limit: int = 100) -> list[BatchItem]:
    """Read Symbol/Ticker and optional Market/Exchange columns from CSV, TXT, or XLSX."""
    ext = os.path.splitext(path)[1].lower()
    if ext in {".xlsx", ".xlsm"}:
        rows = _rows_from_xlsx(path)
    elif ext in {".csv", ".txt", ".tsv"}:
        rows = _rows_from_text(path)
    else:
        raise ValueError("Choose a CSV, TXT, TSV, or XLSX file.")
    if not rows:
        raise ValueError("The selected file is empty.")

    first = [str(v or "").strip().lower() for v in rows[0]]
    symbol_names = {"symbol", "ticker", "stock", "stock symbol", "security"}
    market_names = {"market", "exchange", "country"}
    has_header = any(v in symbol_names for v in first)
    symbol_col = next((i for i, v in enumerate(first) if v in symbol_names), 0)
    market_col: Optional[int] = next((i for i, v in enumerate(first) if v in market_names), None)
    quantity_col: Optional[int] = next((i for i,v in enumerate(first) if v in {"quantity","qty","shares","units"}),None)
    cost_col: Optional[int] = next((i for i,v in enumerate(first) if v in {"cost","cost basis","average cost","avg cost"}),None)
    weight_col: Optional[int] = next((i for i,v in enumerate(first) if v in {"weight","target weight","weight %","allocation"}),None)
    data_rows = rows[1:] if has_header else rows

    items: list[BatchItem] = []
    seen: set[tuple[str, str]] = set()
    for row in data_rows:
        if symbol_col >= len(row):
            continue
        symbol = str(row[symbol_col] or "").strip().upper().replace(" ", "")
        if not symbol or symbol.startswith("#"):
            continue
        raw_market = row[market_col] if market_col is not None and market_col < len(row) else (row[1] if not has_header and len(row) > 1 else "")
        market = _market(raw_market, symbol, default_market)
        key = (symbol, market)
        def cell_number(col):
            if col is None or col>=len(row): return None
            try:
                value=float(str(row[col] or "").replace(",","").replace("%","").strip())
                return value if value>=0 else None
            except (TypeError,ValueError): return None
        if key not in seen:
            seen.add(key); items.append(BatchItem(symbol,market,cell_number(quantity_col),cell_number(cost_col),cell_number(weight_col)))
        if len(items) > limit:
            raise ValueError(f"A maximum of {limit} unique stocks can be analyzed in one batch.")
    if not items:
        raise ValueError("No ticker symbols were found. Use a Symbol column and an optional Market column.")
    return items


def write_csv_row(writer,values) -> None:
    """Prevent spreadsheet formula execution from imported/user-controlled text."""
    safe=[]
    for value in values:
        if isinstance(value,str) and value.startswith(("=","+","-","@")):
            try:
                float(value.replace(",","").removesuffix("%"))
            except ValueError:
                value="'"+value
        safe.append(value)
    writer.writerow(safe)


def export_batch_csv(records: list[tuple[BatchItem, Optional[ValuationResult], str]], path: str) -> None:
    headers=["Input Symbol","Market","Resolved Symbol","Company","Currency","Data Provider","Source Updated","Analysis ID","Parent Analysis ID","App Version","Model Version","Audit Hash","Monte Carlo Seed","Market Price","Fair Value","Upside/Downside %","Required Valuation Buffer %","Valuation Buffer Price","Assessment","Profile","Financial Risk /100","Data Quality /100","Business Quality /100","Governance Risk /100","Order Book /100","Confidence /100","Confidence Label","Reverse DCF Growth %","Monte Carlo P10","Monte Carlo Median","Monte Carlo P90","Probability FV > Price %","Assumption Preset","Manual Override Count","Reviewer Note","Lineage Record Count","Sources Checked","Citation URLs","External Review Count","Live Peers","Base Fair Value","Scenario Weights","Discount Rate %","Cost of Equity %","Beta","Primary Basis","Secondary Model Value","Terminal Value Share %","High-Growth Years","Piotroski F","Altman Z","Beneish M","Median ROIC %","Implied Return %","Error"]
    with open(path,"w",encoding="utf-8-sig",newline="") as handle:
        writer=csv.writer(handle);write_csv_row(writer,headers)
        for item,result,error in records:
            if not result:write_csv_row(writer,[item.symbol,item.market]+[""]*(len(headers)-3)+[error]);continue
            d=result.data;mc=result.monte_carlo_percentiles;capital=result.cost_of_capital or {};forensic=result.forensic_scores or {};piotroski=forensic.get("piotroski");altman=forensic.get("altman");beneish=forensic.get("beneish");secondary=result.scenarios["Base"].graham if "Base" in result.scenarios else None
            engine=[f"{result.base_fair_value:.4f}" if result.base_fair_value is not None else "","/".join(f"{weight:.2f}" for weight in result.scenario_weights),f"{capital['wacc']*100:.2f}" if capital else "",f"{capital['cost_of_equity']*100:.2f}" if capital else "",f"{capital['beta']:.2f}" if capital else "",(result.valuation_drivers or {}).get("basis",""),f"{secondary:.4f}" if secondary else "",f"{result.terminal_value_share*100:.1f}" if result.terminal_value_share is not None else "",result.base_assumptions.high_growth_years,piotroski["score"] if piotroski else "",f"{altman['score']:.2f}" if altman else "",f"{beneish['score']:.2f}" if beneish else "",f"{forensic['roic_median']*100:.1f}" if forensic.get("roic_median") is not None else "",f"{result.implied_return*100:.2f}" if result.implied_return is not None else ""]
            row=[item.symbol,item.market,d.symbol,d.company,d.currency,d.data_provider,d.source_updated_date,result.analysis_id,result.parent_analysis_id,result.app_version,result.model_version,result.audit_hash,result.monte_carlo_seed,f"{d.price:.4f}",f"{result.fair_value:.4f}",f"{result.upside_pct:.2f}",f"{result.dynamic_margin*100:.1f}",f"{result.margin_price:.4f}",result.label,result.methodology_profile,result.risk_score,result.data_quality_score,result.business_quality_score,result.governance_risk_score,result.order_book_score if d.order_book is not None else "",result.confidence_score,result.confidence,f"{result.reverse_dcf_growth*100:.2f}" if result.reverse_dcf_growth is not None else "",f"{mc.get('P10',0):.4f}" if mc else "",f"{mc.get('P50',0):.4f}" if mc else "",f"{mc.get('P90',0):.4f}" if mc else "",f"{result.monte_carlo_probability_above_price*100:.1f}" if result.monte_carlo_probability_above_price is not None else "",result.assumption_preset,len(result.manual_overrides),result.assumption_note,len(d.lineage),"; ".join(d.data_sources),"; ".join(dict.fromkeys(entry.source_url for entry in d.lineage if entry.source_url)),len(d.external_reviews),"; ".join(d.peer_symbols)]+engine+[""]
            write_csv_row(writer,row)
